from scipy import stats

import datetime as dt
import pprint as pp
import numpy as np
import openpyxl
import quandl
import string
import calendar
import sys
import math

import api_key as key

class Stock():
    # initializes the stock with basic data
    def __init__(self, sheet, data):
        self.sheet = sheet
        self.data = data
        self.data_points = len(data)
        self.data_prices = []

        self.bins = 20
        
    # formats the recent data sheet with column titles, etc.
    def format(self):
        # resets formatting
        # not really necessary if the template is clean
        for col in string.ascii_uppercase:
            self.sheet.column_dimensions[col].width = 12
            self.sheet.column_dimensions[col].hidden = False

        self.sheet["A1"] = "Date"
        self.sheet["B1"] = "Close Price"
        
        self.sheet.merge_cells("D1:E1")
        self.sheet["D1"] = "Statistics"
        self.sheet["D2"] = "n"
        self.sheet["D3"] = "mean"
        self.sheet["D4"] = str(self.data_points) + " Day Std Dev"
        self.sheet["D5"] = "40 Day Std Dev"
        self.sheet["D6"] = "20 Day Std Dev"

        self.sheet["D8"] = "BIN AVERAGE"
        self.sheet["E8"] = "FREQUENCY"

    # fills the sheet with data
    def fill_data(self):
        for index in range(0, self.data_points):
            self.data_prices.append(float(self.data[self.data_points - 1 - index][1]))
            self.sheet["A" + str(index + 2)] = str(self.data[self.data_points - 1 - index][0])
            self.sheet["B" + str(index + 2)] = self.data_prices[index]

    # fills the sheet with recent descriptive stats(n, mean, standard deviations for different time periods)
    def fill_stats(self):
        statistics = stats.describe(self.data_prices, bias=False, nan_policy="omit")
        self.sheet["E2"] = statistics[0]
        self.sheet["E3"] = statistics[2]
        self.sheet["E4"] = math.sqrt(statistics[3])
        self.sheet["E5"] = math.sqrt(stats.describe(self.data_prices[self.data_points-40:], bias=False, nan_policy="omit")[3])
        self.sheet["E6"] = math.sqrt(stats.describe(self.data_prices[self.data_points-20:], bias=False, nan_policy="omit")[3])

    # fills the sheet with bin counts and 2 graphs
    def fill_graphs(self):
        # bin size calculations
        minimum, maximum = stats.describe(self.data_prices, bias=False, nan_policy="omit")[1]
        minimum -= .05
        maximum -= .05
        difference = maximum - minimum
        bin_size = difference / self.bins

        # counts for every bin
        counts = [0] * (self.bins + 1)
        for price in self.data_prices:
            counts[int((price - minimum) / bin_size)] += 1

        # puts the bin data into the sheet
        for cell_index in range(0, self.bins):
            bin_name = minimum + (((cell_index * bin_size) + ((cell_index + 1) * bin_size)) / 2)
            self.sheet["D" + str(cell_index + 9)] = bin_name
            self.sheet["E" + str(cell_index + 9)] = counts[cell_index]

        # puts in the bar chart
        bar_chart = openpyxl.chart.BarChart()
        bar_chart.shape = 4
        bar_chart.type = "col"
        bar_chart.style = 10
        bar_chart.title = self.sheet.title + " HISTOGRAM"
        bar_chart.x_axis_title = "BIN AVERAGE"
        bar_chart.y_axis_title = "FREQUENCY"
        bar_data = openpyxl.chart.Reference(self.sheet, min_col = 5, min_row = 8, max_row = 9 + self.bins - 1, max_col = 5)
        bar_categories = openpyxl.chart.Reference(self.sheet, min_col = 4, min_row = 9, max_row = 9 + self.bins - 1)
        bar_chart.add_data(bar_data, titles_from_data = True)
        bar_chart.set_categories(bar_categories)
        self.sheet.add_chart(bar_chart, "G4")

        # puts in the 3 month line chart
        line_chart = openpyxl.chart.LineChart()
        line_chart.style = 12
        line_chart.title = self.sheet.title + " LINECHART"
        line_chart.x_axis_title = "DATE"
        line_chart.y_axis_title = "PRICE"
        line_data = openpyxl.chart.Reference(self.sheet, min_col = 2, min_row = 1, max_row = 1 + len(self.data_prices))
        line_categories = openpyxl.chart.Reference(self.sheet, min_col = 1, min_row = 2, max_row = 1 + len(self.data_prices))
        line_chart.add_data(line_data)
        line_chart.set_categories(line_categories)
        # style the line chart
        style = line_chart.series[0]
        style.graphicalProperties.line.solidFill = "00AAAA"
        style.graphicalProperties.line.dashStyle = "sysDot"
        style.graphicalProperties.line.width = 100050
        self.sheet.add_chart(line_chart, "G22")
